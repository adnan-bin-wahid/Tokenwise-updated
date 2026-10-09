"""Validate revised Word structure, image preservation and rendered page bounds."""

import hashlib
import json
from pathlib import Path
import re
import sys
from zipfile import ZipFile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tmp/report-docx-tools"))
from docx import Document
from docx.oxml.ns import qn
import fitz


def media_hashes(path):
    with ZipFile(path) as package:
        return sorted(hashlib.sha256(package.read(name)).hexdigest()
                      for name in package.namelist() if name.startswith("word/media/"))


def main():
    source = ROOT / "Frineds_report/TokenWise_Final_Report.docx"
    updated = ROOT / "Frineds_report/TokenWise_Final_Report_Updated.docx"
    document = Document(updated)
    paragraphs = [p.text for p in document.paragraphs]
    table_captions = [p for p in paragraphs if re.match(r"Table \d+\.", p)]
    numbers = [int(re.match(r"Table (\d+)\.", p)[1]) for p in table_captions]
    assert numbers == list(range(1, 41)), numbers
    assert media_hashes(source) == media_hashes(updated)
    assert len(media_hashes(updated)) == 20
    chapters = [p.text for p in document.paragraphs if p.style.style_id == "Heading1" and re.match(r"[1-8]\. (?:Project|Requirements|Component|Interface|Testing|User Manual|Validation|Conclusion)", p.text)]
    assert len(chapters) == 8 and chapters[5] == "6. User Manual"
    assert chapters[6] == "7. Validation and Comparative Study"
    assert len(document.tables) == 43
    # Cover table plus the two index tables and forty numbered body tables.
    listed_tables = document.tables[2]
    assert len(listed_tables.rows) == 41
    assert [int(row.cells[0].text) for row in listed_tables.rows[1:]] == list(range(1, 41))
    figure_list = document.tables[1]
    assert [int(row.cells[0].text) for row in figure_list.rows[1:]] == list(range(1, 28))
    assert not any("If you may click" in p or "effectivenes\n" in p or "\ufffd" in p for p in paragraphs)
    all_text = "\n".join(paragraphs) + "\n" + "\n".join(cell.text for t in document.tables for row in t.rows for cell in row.cells)
    for expected in ("1,828", "1,158", "1,511", "36.65%", "353 tokens", "77.78%", "0.3244", "0.012835 g", "0.001746 g", "93.58%", "3.25%", "zero valid native comparison pairs"):
        assert expected in all_text, expected
    worksheet = next(t for t in document.tables if t.rows[0].cells[0].text == "Repository" and t.rows[0].cells[1].text == "Valid pairs /3")
    assert len(worksheet.rows) == 21
    assert all(row.cells[1].text == "Pending" for row in worksheet.rows[1:])
    for paragraph in document.paragraphs:
        if paragraph.text.startswith(("ADD SCREENSHOT", "RESULTS STATUS", "REPLACE AFTER EXECUTION", "FINAL RESULT PARAGRAPH")):
            assert paragraph.style.name == "TW Instruction"
            assert str(paragraph.style.font.color.rgb) == "FF0000"
    assert sum(p.style.name == "TW Instruction" for p in document.paragraphs) == 13
    pdf = fitz.open(ROOT / "tmp/report-review/TokenWise_Final_Report_Updated.pdf")
    outside = []
    empty = []
    locations = {}
    for index, page in enumerate(pdf):
        content = page.get_text()
        if not content.strip() and not page.get_images():
            empty.append(index + 1)
        for block in page.get_text("blocks"):
            if block[0] < -1 or block[1] < -1 or block[2] > page.rect.width + 1 or block[3] > page.rect.height + 1:
                outside.append({"page": index + 1, "bounds": list(block[:4])})
        for term in ("Figure 16.", "Figure 17.", "Figure 18.", "Figure 19.", "Overall live results template", "Per-repository live result worksheet", "7.12 Native Antigravity"):
            if term in content and index > 10:
                locations.setdefault(term, []).append(index + 1)
    assert not outside, outside
    assert not empty, empty
    for pages in locations.values():
        for page in pages:
            pdf[page - 1].get_pixmap(matrix=fitz.Matrix(1, 1)).save(ROOT / f"tmp/report-review/page-{page:03}.png")
    result = {"status": "passed", "checks": ["eight chapters in order", "40 sequential body-table captions", "27 figure-list entries", "20 identical embedded media assets", "13 red instruction paragraphs", "all supplied screenshot values accounted for", "20 pending repository rows without fabricated results", "no blank or off-page rendered content"],
              "pages": len(pdf), "important_pages": locations,
              "source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
              "updated_sha256": hashlib.sha256(updated.read_bytes()).hexdigest()}
    (ROOT / "tmp/report-review/validation.json").write_text(json.dumps(result, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(result, indent=2))


if __name__ == "__main__":
    main()
