"""Produce a revised Word report while preserving the student's original package."""

from copy import deepcopy
import hashlib
import json
from pathlib import Path
import re
import sys
from zipfile import ZipFile

ROOT = Path(__file__).resolve().parents[1]
TOOLS = ROOT / "tmp/report-docx-tools"
if TOOLS.is_dir():
    sys.path.insert(0, str(TOOLS))

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph

SOURCE = ROOT / "Frineds_report/TokenWise_Final_Report.docx"
OUTPUT = ROOT / "Frineds_report/TokenWise_Final_Report_Updated.docx"
NOTES = ROOT / "Frineds_report/TokenWise_Report_Revision_Notes.md"
MANIFEST = ROOT / "demonstration/comparative_study/manifest.json"
PROMPT = "Explain account lockout after failed login attempts and its related tests. Do not modify any files. Cite relevant files and test names."


class Writer:
    def __init__(self, document):
        self.document = document
        self.blocks = []

    def capture(self, element):
        element.getparent().remove(element)
        self.blocks.append(element)

    def p(self, text, style="TW Body"):
        paragraph = self.document.add_paragraph(style=style)
        paragraph.add_run(text)
        self.capture(paragraph._p)
        return paragraph

    def h(self, text, level=2):
        return self.p(text, f"Heading {level}")

    def step(self, number, text):
        return self.p(f"{number}. {text}")

    def code(self, text):
        return self.p(text, "TW Code")

    def instruction(self, text):
        paragraph = self.p(text, "TW Instruction")
        return paragraph

    def image(self, element, caption):
        copy = deepcopy(element)
        # A screenshot can share a paragraph with an old informal caption.
        # Preserve the drawing, not stale prose surrounding it.
        for node in list(copy.iter(qn("w:t"))):
            node.getparent().remove(node)
        # Keep the original relationship and image bytes; only constrain its frame.
        for namespace in ("wp:extent", "a:ext"):
            for extent in copy.findall(f".//{namespace}", {"wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing", "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}):
                cx, cy = int(extent.get("cx", 0)), int(extent.get("cy", 0))
                if cx and cy:
                    scale = min(1, int(Inches(6.5)) / cx, int(Inches(7.5)) / cy)
                    extent.set("cx", str(int(cx * scale)))
                    extent.set("cy", str(int(cy * scale)))
        if copy.tag == qn("w:p"):
            paragraph = Paragraph(copy, self.document._body)
            paragraph.style = self.document.styles["TW Figure"]
            paragraph.paragraph_format.keep_with_next = True
        self.blocks.append(copy)
        self.p(caption, "TW Caption")

    def table(self, title, headers, rows, small=False):
        caption = self.p(f"Table AUTO. {title}.", "TW Caption")
        caption.paragraph_format.keep_with_next = True
        if title.startswith("Per-repository live result worksheet"):
            caption.paragraph_format.page_break_before = True
        table = self.document.add_table(rows=1, cols=len(headers))
        table.autofit = False
        width = 6.5 / len(headers)
        for column in table.columns:
            column.width = Inches(width)
        for cell, text in zip(table.rows[0].cells, headers):
            cell.text = str(text)
        header_props = table.rows[0]._tr.get_or_add_trPr()
        repeat = OxmlElement("w:tblHeader")
        header_props.append(repeat)
        for row in rows:
            cells = table.add_row().cells
            for cell, text in zip(cells, row):
                cell.text = str(text)
        borders = OxmlElement("w:tblBorders")
        for name in ("top", "left", "bottom", "right", "insideH", "insideV"):
            edge = OxmlElement(f"w:{name}")
            edge.set(qn("w:val"), "single")
            edge.set(qn("w:sz"), "4")
            edge.set(qn("w:color"), "B7C1C9")
            borders.append(edge)
        table._tbl.tblPr.append(borders)
        for index, row in enumerate(table.rows):
            props = row._tr.get_or_add_trPr()
            props.append(OxmlElement("w:cantSplit"))
            for cell in row.cells:
                if index == 0:
                    shade = OxmlElement("w:shd")
                    shade.set(qn("w:fill"), "E8EDF1")
                    cell._tc.get_or_add_tcPr().append(shade)
                for paragraph in cell.paragraphs:
                    paragraph.style = self.document.styles["TW Body"]
                    paragraph.paragraph_format.space_after = Pt(1 if small else 3)
                    paragraph.paragraph_format.space_before = Pt(1 if small else 3)
                    if small:
                        paragraph.paragraph_format.line_spacing = 1.0
                    paragraph.paragraph_format.keep_with_next = index == 0
                    for run in paragraph.runs:
                        run.font.size = Pt(8 if small else 9)
                        run.bold = index == 0
        self.capture(table._tbl)
        return table


def configure_styles(document):
    definitions = {
        "TW Body": ("Times New Roman", 11, 6), "TW Caption": ("Times New Roman", 10, 6),
        "TW Instruction": ("Times New Roman", 10, 6), "TW Code": ("Consolas", 9, 6),
        "TW Figure": ("Times New Roman", 10, 3),
    }
    for name, (font, size, after) in definitions.items():
        style = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = document.styles["normal"]
        style.font.name = font
        style.font.size = Pt(size)
        style.paragraph_format.line_spacing = 1.15
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_together = True
        if name == "TW Caption":
            style.font.italic = True
        if name == "TW Instruction":
            style.font.color.rgb = RGBColor.from_string("FF0000")
        if name == "TW Figure":
            style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER


def text(element):
    return "".join(node.text or "" for node in element.iter(qn("w:t")))


def replace_text(element, replacement):
    paragraph = Paragraph(element, None)
    properties = deepcopy(paragraph.runs[0]._r.rPr) if paragraph.runs and paragraph.runs[0]._r.rPr is not None else None
    paragraph.clear()
    run = paragraph.add_run(replacement)
    if properties is not None:
        run._r.insert(0, properties)


def build_manual(document, original):
    w = Writer(document)
    w.h("6. User Manual", 1)
    w.h("6.1 Prerequisites and Installation Files")
    w.p("The normal user installs TokenWise as a VSIX extension in Antigravity. The procedure does not require cloning the source repository, compiling TypeScript, pressing F5, or starting a development server. This manual describes version 0.6.8 and the single demonstration/tokenwise_demo application. An arbitrary trusted local Python repository can be used instead.")
    w.p("Figure 7 summarizes the prerequisites: Antigravity with workspace rules and a command tool; 64-bit Python 3.12; a local Python folder; initial dependency/model download access; approximately 10 GB free disk; and adequate memory. Eight GB RAM is a starting recommendation, not a guarantee for long inputs. CPU operation is supported; a GPU, Ollama and an MCP server are not required. Antigravity access and billing are separate from TokenWise, which needs no API key of its own. Windows is the verified native platform; portable launchers do not establish completed macOS/Linux testing.")
    w.image(original[289], "Figure 7. TokenWise prerequisite and environment information on the distribution website.")
    w.p("The website's statement that retrieval runs locally refers to TokenWise indexing and pruning. Once context is supplied to Antigravity, the configured downstream model may process it externally. The prerequisite illustration is guidance, not evidence that installation or inference completed.")
    w.h("6.2 Download, Install and Upgrade")
    w.step(1, "Open https://tokenwise-extension.vercel.app/ and select Download for Antigravity (.vsix). Alternatively, use the versioned v0.6.8 GitHub release linked in the project README. The demonstration ZIP contains presentation materials; GitHub's source-code ZIP is not the extension installer.")
    w.step(2, "In Antigravity, open Extensions, select the three-dot menu and choose Install from VSIX. Select tokenwise-vscode-0.6.8.vsix and reload the window when requested.")
    w.step(3, "Open the command palette with Ctrl+Shift+P and run TokenWise: Set Up Backend. When updating, update the backend as well as the VSIX so that the service supports the same features.")
    w.step(4, "Open the intended Python folder, run TokenWise: Enable Automatic Context and start a new chat after configuration. Re-enable after an upgrade to refresh owned rules and launchers.")
    w.p("Figure 8 shows the distribution page with separate VSIX, demonstration-bundle and repository buttons. Its four installation panels correspond to downloading the package, opening Extensions, installing from VSIX and reloading the editor. The web page distributes the extension; it is not the runtime repository-pruning service.")
    w.image(original[298], "Figure 8. TokenWise download page and the four-step VSIX installation guide.")
    w.h("6.3 Managed Backend Setup and Recovery")
    w.p("On a new computer, choose Install Managed Backend, review the confirmation and allow setup to finish. The installer prepares a private environment and integrity-checked runtime shared by configured workspaces. If a complete development backend already exists, Use Existing Backend is an alternative: select its root, not the target application or the inner swe-pruner folder.")
    w.p("Figure 9 explains the seven stages. It is a website illustration of the process, not a successful installation log. Stage 7 registers the validated backend; workspace rules and launchers are then configured by Enable Automatic Context. These are related operations rather than evidence that every workspace is automatically configured by backend registration.")
    w.image(original[301], "Figure 9. Illustration of the seven managed backend installation stages.")
    w.table("Installation stages and retry actions", ["Stage", "Operation", "If the stage fails"], [
        ["1. Prerequisites", "Find 64-bit Python 3.12 and verify bundled files", "Install or repair Python 3.12; restart Antigravity; set TokenWise Python Path if detection fails."],
        ["2. Backend files", "Copy verified runtime into private storage", "Check free space and write permissions; retry the failed step."],
        ["3. Environment", "Create or verify the private Python environment", "Correct the interpreter/environment error; retain validated model caches."],
        ["4. Dependencies", "Install packaging tools, CPU PyTorch and backend dependencies", "Correct network/proxy/package-host access; retry and reuse validated substeps."],
        ["5. Model", "Download approximately 1.35 GB of pinned weights and verify SHA-256", "Check model-host access and storage; allow supported resumption and checksum verification."],
        ["6. Verification", "Check imports, tokenizer and carbon-estimation artifacts", "Read the exact import/artifact error in TokenWise Setup and fix its cause."],
        ["7. Registration", "Save the verified installation and its user setting", "Correct storage/settings permissions and retry without deleting verified weights."],
    ])
    w.p("Choose View > Output > TokenWise Setup for the detailed error. Use Retry Failed Step if offered. If the notification was dismissed or setup was cancelled, run Set Up Backend again in the same editor profile. Do not uninstall or erase caches as the first recovery action. Setup completion time depends on the network and hardware and is not an end-to-end response benchmark.")
    w.p("Figure 10 is an actual editor notification at Step 3/7: Creating a private Python environment. The Source label identifies TokenWise and Cancel stops the current setup. It shows progress at that stage, not completion of all seven stages.")
    w.image(original[319], "Figure 10. Backend setup notification during private environment creation, Step 3/7.")
    w.h("6.4 Verify Installation and Enable the Workspace")
    w.step(1, "Use File > Open Folder to open demonstration/tokenwise_demo, or the root of another local Python application. Trust only a repository whose contents you trust.")
    w.step(2, "Run TokenWise: Enable Automatic Context and confirm the selected workspace. In a multi-folder window, enable each required repository separately.")
    w.step(3, "Open Agent Settings > Customizations > Rules and verify the TokenWise rule. Existing unrelated rules and hooks should remain unchanged.")
    w.step(4, "Run TokenWise: Diagnose Setup and resolve any reported error. Check Backend Health or Start Backend can be used for explicit inspection; normal retrieval starts the registered backend when needed.")
    w.p("Setup adds owned integration files under .agents and a backend link/activity directory under .tokenwise. It does not require the application to be inside a TokenWise checkout. Unrelated customization is preserved; malformed JSON and customized owned launchers are not silently replaced.")
    w.p("Figure 11 shows TokenWise in the Extensions view with its icon, publisher adnan-bin-wahid, description and Disable/Uninstall controls. Those controls indicate that the extension is installed. The details describe conversation-memory behavior for 0.6.8. The image does not itself prove backend readiness or a public marketplace listing; installing a VSIX also creates an installed entry.")
    w.image(original[321], "Figure 11. Installed TokenWise extension and its available lifecycle controls.")
    w.h("6.5 Ask an Ordinary Antigravity Question")
    w.step(1, "Start a fresh Antigravity chat. Close unnecessary source editors, clear highlighted code and attach no files when demonstrating automatic discovery.")
    w.step(2, "Send the following read-only question. The same exact wording is used for the comparison in Chapter 7.")
    w.code(PROMPT)
    w.step(3, "Allow the local TokenWise launcher if the editor asks for permission. Wait for its completed context output; do not interpret a waiting timer as successful retrieval.")
    w.step(4, "Inspect the final answer and run TokenWise: Show Automatic Context. Match the current prompt, event/timestamp and included excerpts. A status-bar entry or old latest.json alone is not proof of delivery for this request.")
    w.p("Figure 12 records a different illustrative prompt, Explain me the project. The agent waits for TokenWise and then reports context from app.py, reports.py, security/auth_service.py, security/models.py, security/settings.py and workflows.py before displaying a project summary. This shows the agent's visible invocation/acknowledgement in that conversation. Its waiting messages and Worked durations must not be treated as the total time of a controlled comparison. They also do not reveal billed tokens or all internal file reads.")
    w.image(original[324], "Figure 12. Recorded Antigravity project-overview conversation with TokenWise retrieval acknowledgement.")
    w.p("Automatic context is the normal prompt workflow. The teaching commands below prepare inspectable outputs independently; running them does not by itself inject their output into an Antigravity conversation.")
    w.h("6.6 Discover the Extension Commands")
    w.p("Press Ctrl+Shift+P and type TokenWise. Figure 13 shows the command palette with setup, automatic context, pruning, comparison and guide commands. The highlighted Demonstrate Pruning Inputs command is for inspecting the three input modes; it is not required before every chat prompt.")
    w.image(original[327], "Figure 13. TokenWise commands available through the Antigravity command palette.")
    w.table("Common commands and when to use them", ["Command, prefixed TokenWise:", "Purpose"], [
        ["Set Up Backend / Diagnose Setup", "Install/update the runtime and inspect failures or readiness."],
        ["Enable Automatic Context / Show Automatic Context", "Configure the current workspace and inspect the latest prepared result."],
        ["Demonstrate Pruning Inputs", "Choose repository discovery, selected source or controlled conversation replay."],
        ["Prune Current File / Prune Selected Code", "Prune the whole active file or exactly the highlighted source."],
        ["Build Repository Context", "Prepare a manual task-aware workspace packet."],
        ["Compare Context Strategies", "Compare explicit context-packet strategies using the local tokenizer."],
        ["Configure Automatic Comparison", "Enable, disable or retry background prepared-packet comparison."],
        ["Import Antigravity Usage Comparison", "Import two genuine independent successful CLI logs; no model call is launched."],
        ["Open Setup Guide / Open Demonstration Guide / Open Validation Guide", "Open bundled instructions for recovery or presentation."],
    ])
    w.h("6.7 Repository Discovery without a Selected File")
    w.step(1, "Run Demonstrate Pruning Inputs and choose Repository: no selected file.")
    w.step(2, "In the subsequent task input, enter: Explain session expiry and revocation, not invoice pricing.")
    w.step(3, "Enter 0.45 in the threshold input and confirm. Wait for the Repository Context result.")
    w.p("Figure 14 is the mode-selection menu, not the task input. Its blank search box filters the three choices: Repository: no selected file, Selected file or highlighted excerpt, and Conversation: replay earlier user context. The task question is entered after choosing a mode. Repository mode intentionally supplies no active-file hint even if an editor happens to be open.")
    w.image(original[328], "Figure 14. Three input scopes offered by Demonstrate Pruning Inputs.")
    w.p("Figure 15 shows the separate threshold input. A value must be between 0 and 1; Enter confirms and Escape cancels. The requested 0.45 is the user setting, not a promised 45% saving. The repository policy can apply tier-adjusted thresholds or use a non-neural representation, as the included-file table later discloses.")
    w.image(original[330], "Figure 15. Requested pruning-threshold input set to 0.45.")
    w.h("6.8 Read the Repository Context View")
    w.p("Figures 16 and 17 are two portions of one recorded repository-discovery result. They illustrate the session-expiry/revocation task rather than the account-lockout comparison. Values are transcribed from the supplied images, not newly measured during this report revision. The task's invoice exclusion guides source selection but does not guarantee that every retained supporting declaration is unrelated to billing.")
    w.image(original[332], "Figure 16. Repository discovery input trace, memory state and synthesized goal.")
    w.table("Repository input-trace and goal fields", ["Visible field", "Recorded value or state", "Interpretation"], [
        ["Mode", "Repository discovery (no file selected)", "Search starts from the prompt rather than highlighted source."],
        ["Current task", "Explain session expiry and revocation, not invoice pricing", "The latest request, including its exclusion."],
        ["Code scope", "Repository discovery without editor hints", "No active-file anchor is supplied by this teaching mode."],
        ["Requested threshold", "0.45", "Requested starting threshold; effective file thresholds may differ."],
        ["History source", "None", "No earlier user reference was supplied for this run."],
        ["Indexed Python files", "11", "Eligible Python files in the indexed workspace; not all are sent."],
        ["Inference objective", "Expandable", "The effective task-conditioned objective used for retrieval/pruning."],
        ["Conversation memory", "0 selected / 0 considered; omitted 0; 0 characters", "No eligible earlier user turns were considered or injected."],
        ["Task reference / Earlier requirements", "none / none", "No historical topic or retained requirement is reported."],
        ["Outgoing reference", "not_used / 0 tokens", "There is no outgoing memory block in this packet."],
        ["Export Pruning Run", "Action button", "Save input scope and result evidence for inspection."],
        ["Task type", "generic_task", "Deterministic classification of this explanation request; not an accuracy label."],
        ["Objective / Identifiers", "Task-focused objective; session, expiry, revocation", "Converted task focus and search identifiers. Excluded topics and diagnostics, when shown, describe constraints/errors, not successful test execution."],
    ])
    w.image(original[333], "Figure 17. Repository source/packet metrics, configured carbon estimates and included-file provenance.")
    w.table("Repository result token values in Figure 17", ["Field", "Value", "Meaning"], [
        ["Source tokens", "1,828", "Original source token counts for the included files, not the entire repository or provider request."],
        ["Retained source", "1,158", "Source remaining after the selected pruning/representation policy."],
        ["Packed tokens", "1,511", "Complete outgoing reference packet, including its formatting and other bounded blocks."],
        ["Source reduction", "36.65%", "100 x (1,828 - 1,158) / 1,828; a source-only comparison."],
        ["Formatting overhead", "353 tokens", "Reported packet overhead: 1,511 - 1,158. It is not extra source."],
        ["Files included", "8", "Files represented in this packet, out of 11 indexed Python files."],
    ])
    w.p("The source reduction must not be calculated by comparing 1,828 source tokens directly with a differently formatted full packet. Packed tokens are the relevant size for TokenWise's complete-packet budget, but remain local-tokenizer counts. Guidance and memory consume packet space when present. If an automatic result also shows Prepared in, that is TokenWise preparation time, not total Antigravity answer time.")
    w.table("Included-file values transcribed from Figure 17", ["File", "Relation / tier", "Source / retained", "Score", "Method / applied threshold"], [
        ["tests/test_workflows.py", "related test / 1", "421 / 203", "0.9964", "scope_filter+neural_lines / 0.30"],
        ["app.py", "direct/relevant dependency / 2", "313 / 269", "0.9987", "scope_filter+neural_lines / 0.60"],
        ["tests/test_auth.py", "related test / 2", "410 / 182", "0.1755", "neural_lines / 0.60"],
        ["workflows.py", "small task-matched source / 2", "274 / 177", "0.1608", "scope_filter+short_source_retained / not applied"],
        ["security/models.py", "small task-matched source / 2", "117 / 117", "0.0919", "short_source_retained / not applied"],
        ["tests/test_models.py", "related test / 2", "144 / 144", "0.0919", "short_source_retained / not applied"],
        ["security/settings.py", "dependency interface / 3", "23 / 22", "0.0230", "signature_interface / not applied"],
        ["reports.py", "dependency interface / 3", "126 / 44", "0.0000", "signature_interface / not applied"],
    ], small=True)
    w.p("File is the repository-relative source path; Relation explains its discovery or representation; Tier controls treatment priority. Source and Retained are per-file local counts, whose totals are 1,828 and 1,158. Score is a retrieval/ranking signal, not a probability that the final answer is correct. A file can remain as supporting context despite a low score. Tier 1 uses lighter pruning here (0.30), while neural Tier 2 uses 0.60. The requested 0.45 is therefore not applied identically to every file. Not applied means that this row used scope filtering, short-source retention or an interface rather than a neural threshold. The final Unified context is the actual excerpt to inspect; filenames alone do not establish preserved behavior. Copy Unified Context copies that packet, not an answer or the entire repository.")
    w.h("6.9 Interpret the Carbon and Energy Values")
    w.p("The Carbon impact region in Figure 17 compares the same included files and formatting without pruning against the prepared packet under a fixed inference scenario. It does not compare a live disabled Antigravity conversation with an enabled one, and does not compare the entire repository against eight files.")
    w.table("Configured repository carbon and energy values in Figure 17", ["Quantity", "Before", "After / difference"], [
        ["CO2", "0.171382 g", "0.158548 g"],
        ["Prefill energy", "316.9735 J", "219.7005 J"],
        ["Decode energy", "981.9252 J", "981.9252 J"],
        ["Total energy", "1,298.8988 J", "1,201.6257 J"],
        ["Energy saved", "Before minus after", "97.2731 J"],
        ["CO2 saved", "Before minus after", "0.012835 g"],
    ])
    w.p("Prefill estimates processing the supplied input; Decode estimates generating the configured expected output. Decode is unchanged because the output-length assumption is held constant. Total is the sum of the two phases. Energy saved and CO2 saved are signed differences calculated from full-precision predictions; subtracting the rounded displayed endpoints can differ in the final decimal. Smaller input therefore does not imply the same percentage reduction in total energy.")
    w.table("Estimator assumptions and provenance shown in the result", ["Field", "Recorded value", "Explanation"], [
        ["Model", "meta-llama-3-8b-instruct", "Estimator target; not automatically the model selected in Antigravity."],
        ["Prefill route / Decode route", "xgboost_interpolation / xgboost_interpolation", "Trained artifact family used for each phase within the configured size range."],
        ["Feature source", "artifact_models:model_registry+request_overrides", "Registry-backed model/hardware attributes combined with explicit request assumptions."],
        ["Carbon intensity", "475.00 gCO2/kWh", "Assumed electricity emissions factor; not a measurement of the user's or provider's power grid."],
        ["Baseline", "Same files and formatting, without pruning", "Matched formatted-context comparison, distinct from all-code or native-agent baselines."],
        ["Disclaimer", "Approximate inference estimates, not measured emissions", "Predictions exclude local TokenWise work and do not establish net environmental savings."],
    ])
    w.p("The conversion is CO2_g = energy_J / 3,600,000 x intensity_g_per_kWh. For example, 97.2731 J at 475 gCO2/kWh is approximately 0.012835 g. A pending or unavailable estimate must remain pending/unavailable, not zero. Positive expansion can appear as Energy increase or CO2 increase. The estimator's expected output-token setting should be recorded in the export even when it is not visible in the screenshot.")
    w.h("6.10 Prune a Selected File or Highlighted Excerpt")
    w.step(1, "Open security/models.py. Clear the selection, run Demonstrate Pruning Inputs and choose Selected file or highlighted excerpt.")
    w.step(2, "Enter Explain session class and threshold 0.45 to reproduce the input represented in Figure 18. Prune Current File is also a direct whole-file command.")
    w.step(3, "For excerpt mode, highlight the Session decorator and class body and repeat the teaching command. The captured scope changes to Selected excerpt; unselected Account code was never input in that run.")
    w.image(original[335], "Figure 18. Selected-file neural pruning of security/models.py for the Session class.")
    w.p("Figure 18 reports Selected entire file, the task Explain session class, the full security/models.py path, requested threshold 0.45, History source None and First source line 1. Inference objective expands the actual goal; Export Pruning Run saves the trace. These fields identify the input before judging the output.")
    w.table("Neural-pruning values shown in Figure 18", ["Field", "Value", "Interpretation"], [
        ["Relevance", "0.3244", "Document-level task-relevance signal; not answer correctness or a percentage of lines retained."],
        ["Original", "117 tokens", "Whole selected file counted by TokenWise's tokenizer."],
        ["Pruned", "26 tokens", "Emitted pruned text, including its elision formatting."],
        ["Reduction", "77.78%", "100 x (117 - 26) / 117; not a provider-usage measurement."],
    ])
    w.table("Selected-file energy and carbon values in Figure 18", ["Quantity", "Before", "After / difference"], [
        ["CO2", "0.131804 g", "0.130058 g"],
        ["Prefill energy", "17.0119 J", "3.7804 J"],
        ["Decode energy", "981.9252 J", "981.9252 J"],
        ["Total energy", "998.9371 J", "985.7057 J"],
        ["Energy saved", "Before minus after", "13.2315 J"],
        ["CO2 saved", "Before minus after", "0.001746 g"],
    ])
    w.p("The model, routes, feature source and 475.00 gCO2/kWh factor match Figure 17, but Baseline is Source text only: it compares this selected source before and after pruning rather than full formatted repository context. The 77.78% text reduction is accompanied by only about 1.32% predicted total-energy reduction because the decode workload stays fixed.")
    w.p("Below the visible region, inspect Pruning details, Line decisions, Original and Pruned context. Model input tokens include task/instruction overhead and differ from source tokens. Kept line fragments identify the decision mask; each Source line is an original line number, Mean relevance is its aggregated score, and Decision mask distinguishes Keep: threshold, Keep: preservation/gap and Not selected by mask. Unscored lines are not relevance zero. Formatting can restore tiny gaps, so the emitted text is the final reference. A retained Session body may omit @dataclass or imports: inspect originals before explaining constructor behavior or editing. Copy Pruned Context copies the excerpt. Filtered-line markers are not executable Python.")
    w.h("6.11 Demonstrate Same-Conversation Context")
    w.step(1, "In one new chat with automatic context enabled, send the account-lockout prompt from Section 6.5 and wait for the final answer.")
    w.step(2, "In that same chat, send: What about its expiry boundary? Do not modify any files.")
    w.step(3, "Show Automatic Context and inspect the new task, effective objective, History source and Conversation memory. Confirm that the subject is account lockout rather than unrelated session or invoice behavior.")
    w.p("Figure 19 shows the preceding lockout question, a reference to test_wrong_password_after_expiry_starts_a_new_count, and the follow-up about its expiry boundary. The visible launcher command reformulates the subject as account lockout expiry boundary and related tests and encodes it with QueryBase64. Encoding transports the query; it is not token reduction or encryption. Working indicates that the command is still running. This screenshot demonstrates visible topic continuity and invocation, but does not display the completed follow-up answer, the bounded memory inspector or provider usage.")
    w.image(original[337], "Figure 19. Same-chat account-lockout follow-up and visible subject-focused TokenWise launcher invocation.")
    w.p("The memory inspector separately reports selected/considered turn counts, omissions, reference characters, reasons, task reference, recognized earlier requirements and outgoing status/token count. Selected history is not necessarily all injected: the packet can include a full or compact reference block or omit it to preserve evidence space. Native scoped user turns and agent-supplied fallback references have different labels; a controlled replay is not native transcript capture. Assistant answers and tool output are excluded as user requirements. Selection is bounded to eight relevant earlier turns and 4,000 characters and uses lexical heuristics, not unlimited or perfect semantic memory.")
    w.instruction("ADD SCREENSHOT - Figure 20: After the follow-up completes, capture Show Automatic Context with the current task, expanded Conversation memory, selected messages/reasons, history source and exact outgoing status/token count. Caption: Inspectable bounded user references for a completed follow-up. Use the actual capture date.")
    w.p("For a separate teaching replay, choose Conversation: replay earlier user context, supply the earlier lockout task and ask Which tests cover that behavior? Label it replay. For a new-chat control, ask only that vague question in a fresh chat; the extension must not invent the prior conversation. Memory can be disabled with conversation_memory: false in .agents/tokenwise.json and restored afterward. Keep these demonstrations outside independently timed WITH/WITHOUT trials.")
    w.h("6.12 Inspect Prompt Guidance, Comparison and Export")
    w.p("In a repository result, inspect Response guidance and the Unified context. The outgoing task-aware guidance asks for grounded claims, real file/symbol citations, respect for user constraints and explicit uncertainty. Its version, task profile, full/compact/omitted status and local token cost are inspectable. It is separate from source data and adds no extra goal-generation model call by default. Guidance is intended to support better answers; its presence is not proof that answer accuracy improved.")
    w.instruction("ADD SCREENSHOT - Figure 21: Capture the actual Response guidance metadata and its outgoing block, alongside the source-reference boundary. Caption: Bounded task-aware prompt guidance in the prepared context.")
    w.step(1, "Run Configure Automatic Comparison and choose Enable automatic packet comparison. Send a fresh ordinary prompt with TokenWise enabled.")
    w.step(2, "Wait for Context strategy comparison in Show Automatic Context. Inspect All-Python baseline, TokenWise packet, packet change, task-plus-packet counts, file lists and estimator assumptions. Export Comparison saves the evidence; a stale-snapshot error requires fresh retrieval.")
    w.step(3, "Disable this background comparison before timed native trials. The extra local comparison is an inspection feature, not necessary for ordinary context delivery.")
    w.p("The comparison assumes all indexed Python code is supplied in its baseline. Source tokens count source; Packet tokens include packet formatting; Task + packet tokens also include the task. The displayed change can be a genuine increase for small inputs. Energy/CO2 columns remain estimates. This is not automatic monitoring of Antigravity's file reads, billing or actual model requests.")
    w.instruction("ADD SCREENSHOT - Figure 22: Capture a completed automatic packet comparison with exact task, both file lists, locally counted complete packets and the signed change. Caption: Hypothetical all-Python versus TokenWise prepared-packet comparison. Do not label it actual Antigravity usage.")
    w.p("For real reported agent usage, run Import Antigravity Usage Comparison and select two independent successful single-turn CLI JSON/stream-JSON logs from matched WITHOUT and WITH runs. The importer validates sessions, available model identity, usage, answer and completed tools. It does not start paid calls or grade answer quality. A manual IDE answer may have no reported input counter; record N/A rather than substituting the local packet count. Export Pruning Run, Export Comparison and saved chat/tool evidence answer different reproducibility questions.")
    w.h("6.13 Troubleshooting, Readiness and Cleanup")
    w.table("User recovery actions", ["Symptom", "Action and verification"], [
        ["No automatic context", "Check trusted local workspace, enabled setting, active TokenWise rule, selected registered backend and command permission. Send a fresh prompt after correction."],
        ["Backend unreachable / older feature missing", "Run Diagnose Setup, Set Up Backend and Start Backend as applicable; refresh workspace rules after upgrading."],
        ["Permission denied", "Review and approve only the required local TokenWise operation. Do not enable blanket command bypass or report a denied run as successful."],
        ["Old result shown", "Match query/event/timestamp and send a new prompt. Close the old panel; a previous latest.json does not prove current activity."],
        ["CO2 pending, disabled or unavailable", "Check enableCarbonEstimation and diagnostic/artifact errors. Retrieval can succeed independently; missing estimation is not zero emissions."],
        ["Large input is slow", "Allow cold startup, inspect timeouts and source scope, and test a focused excerpt. A small output budget does not itself cap neural input computation."],
        ["Comparison stale or unavailable", "Retrieve again on the current saved snapshot and retry through Configure Automatic Comparison."],
        ["No reported provider tokens", "Keep actual input tokens N/A. Store local packet tokens in a separately labeled field."],
    ])
    w.instruction("ADD SCREENSHOT - Figure 23: Capture actual Diagnose Setup output showing the selected backend, readiness and any resolved error. Caption: Verified backend readiness before live validation. Do not replace this with a website setup illustration.")
    w.p("For temporary disabling, set enabled: false in .agents/tokenwise.json and change the TokenWise rule to Manual without invoking it. A clean comparison baseline also requires no active duplicate rule or injected packet. Preserve backups and restore the original rule/settings afterward. For removal, TokenWise: Remove All Local Data and extension uninstall use ownership-aware cleanup. Modified, shared or unverifiable resources can be retained with warnings; never promise deletion of arbitrary user files or shared caches.")
    return w.blocks


def build_native_comparison(document):
    w = Writer(document)
    w.h("7.12 Native Antigravity WITH/WITHOUT Evaluation")
    w.p("This evaluation compares Antigravity's normal repository-assisted answer with the same workflow using TokenWise. It is distinct from the completed local packet study in Sections 7.3-7.11. WITHOUT retains native file-reading tools; it is not an all-code packet forced into a prompt. WITH must demonstrate delivery of a fresh TokenWise packet, not just an installed extension or an old result.")
    w.instruction("RESULTS STATUS - The live comparison tables below are fill-ready templates, not measured outcomes. Replace bracketed fields only after saving the real runs. The 120-run evaluation has not been completed during report revision; do not submit Pending fields as experimental results.")
    w.p("The planned scope is 20 pinned Python repositories with one predefined task each, three paired repetitions per repository and two independent conditions per pair: 60 pairs and 120 runs. The numbered source folders are under demonstration/comparative_study. Exact tasks are in TASKS.md, provenance in manifest.json and the 120-row capture worksheet in results-template.csv. The folders are source snapshots, not individual Git clones. Keep searches inside the opened numbered folder; do not let Git commands search the parent TokenWise repository.")
    w.p("Before the full study, rehearse one lockout pair in demonstration/tokenwise_demo using the following question. This is a separate functional pilot, not one of the 20 upstream repository tasks. Each formal repository uses its own source-grounded six-fact rubric defined before either answer is inspected; the lockout rubric cannot grade Click or another unrelated task.")
    w.code(PROMPT)
    w.table("Live comparison settings to record before execution", ["Control", "Recorded setting"], [
        ["Actual evaluation date(s)", "[INSERT actual dates]"],
        ["Editor/CLI version and transport", "[INSERT IDE or CLI; do not combine their telemetry]"],
        ["TokenWise/backend revision", "[INSERT installed version and backend revision]"],
        ["Pinned model / reasoning effort", "[INSERT identical model and effort for both conditions]"],
        ["Task, source snapshot and rubric", "[INSERT task IDs, full commits and frozen six-fact rubric evidence]"],
        ["Token budget / threshold / candidate limit", "[INSERT actual automatic settings; fresh-workspace defaults are 4096 / 0.45 / 6]"],
        ["Native tools and permissions", "[INSERT matched usable policy; no deliberately restricted baseline]"],
        ["Background packet comparison", "Disabled during timed runs"],
        ["Backend state and initialization", "[INSERT warm/cold state and any pre-run warm-up]"],
        ["Attached files, selection and previous history", "None; each run begins in a fresh independent chat"],
    ])
    w.h("7.12.1 WITHOUT Condition", 3)
    w.step(1, "Back up .agents/tokenwise.json and the TokenWise rule outside the active project's context. Set the existing enabled flag to false. Set only the TokenWise rule activation to Manual and do not invoke it; verify no duplicate active rule or hook injects context.")
    w.step(2, "Disable automatic packet comparison, close prior result/source panels, clear highlighted code and attach no files. Leave Antigravity's native reading tools available.")
    w.step(3, "Start a new chat, start a stopwatch when submitting the fixed task and stop it when the final answer finishes. Save the entire answer and visible tool activity or a genuine CLI trace.")
    w.step(4, "Verify no new TokenWise invocation or supplied packet occurred. Record model, effort, source snapshot, conversation identity, timing, observed calls, actual reported counters or N/A, and failure status.")
    w.instruction("ADD SCREENSHOT - Figure 25: After a genuine successful WITHOUT run, capture the fixed task, final answer, selected model and expanded native reading tools, with the disabled integration recorded separately. Caption: Native Antigravity answer without TokenWise. Do not use the permission-blocked preflight as a successful baseline.")
    w.h("7.12.2 WITH Condition", 3)
    w.step(1, "Restore enabled: true and the TokenWise rule's original activation, normally Always On. Keep source, model, reasoning effort, permissions and editor state matched.")
    w.step(2, "Start another fresh chat and submit the identical task. Time submission through final completion, including TokenWise retrieval and agent tools.")
    w.step(3, "Save the answer, visible context delivery/tool trace, fresh task/event/timestamp and actual supplied packet. Inspect its excerpts and evidence, not only its included-file names.")
    w.step(4, "Record total answer time and TokenWise preparation time separately. Keep provider input/output/cache/thinking counters separate from locally counted context tokens. Verify source bytes are unchanged.")
    w.instruction("ADD SCREENSHOT - Figure 26: Capture the independent matched WITH answer and completed TokenWise invocation/output, together with the current task and model. Caption: Native Antigravity answer with fresh TokenWise context. A status-bar count alone is insufficient.")
    w.p("Alternate order across the three repetitions: WITHOUT then WITH; WITH then WITHOUT; WITHOUT then WITH. Record whether the backend was cold or warm rather than attributing startup differences to answer quality. Any permission denial, timeout, quota failure, empty answer, source modification or baseline contamination remains a failed/invalid trial. Preserve it before a replacement run; do not silently substitute zero-valued metrics.")
    w.h("7.13 Source-Grounded Answer Scoring")
    w.p("For the lockout pilot, the original settings, service and tests provide the six facts below. AuthService.authenticate checks now < locked_until before accepting a password. At an existing deadline it clears the old lockout state before the next password check. The fixture uses integer time, dataclasses and teaching-only password hashing; it has no background unlock thread, JWT subsystem or persistent user database.")
    w.table("Six-fact rubric for the separate account-lockout pilot", ["ID", "Supported fact", "Verification source"], [
        ["F1", "The third failed password attempt triggers lockout", "security/settings.py; test_lockout_threshold"],
        ["F2", "The deadline is the triggering failure time plus 60 seconds", "security/settings.py; AuthService.authenticate; test_lockout_threshold"],
        ["F3", "Even a correct password is rejected before the deadline", "test_correct_password_rejected_during_lockout"],
        ["F4", "At the exact deadline a correct password succeeds and both failure/deadline fields reset", "test_expiry_boundary"],
        ["F5", "A successful login before lockout resets accumulated failures", "test_success_resets_failures_before_lockout"],
        ["F6", "A wrong password at expiry clears the old lock and starts a new failure count of one", "test_wrong_password_after_expiry_starts_a_new_count"],
    ])
    w.p("All named tests are in demonstration/tokenwise_demo/tests/test_auth.py. Three wrong attempts at time 100 produce a deadline of 160: a correct password fails at 159 and succeeds at 160. A wrong password at expiry yields False, failed_attempts = 1 and locked_until = 0. Define equivalent task-specific rubrics from original source/tests for all 20 upstream tasks before scoring them.")
    w.p("Award one point for each accurately stated, source-supported rubric fact, maximum six. A missing or incorrect fact earns zero for that criterion, but is not automatically an unsupported claim. Count distinct false factual assertions as incorrect and distinct assertions lacking repository evidence as unsupported; do not count the same assertion twice or count repeated wording multiple times. Keep the two categories separately in the worksheet and use their sum only in the combined summary row. Check exact cited functions/test names. Token-saving claims and generic recommendations do not replace behavior facts. Student scoring should retain supporting notes and can be reviewed by the supervisor; it is not automatically blinded independent assessment.")
    w.h("7.14 Measurement and Results Worksheets")
    w.p("One run record is required for each of the 120 planned runs. Its fields include repository/commit, task, repeat, condition, chat ID, model/effort, backend state, answer, rubric marks, incorrect and unsupported claim notes, total wall time, preparation time, tool evidence, reported usage source and success/failure status. One paired record links two independent chats with identical controls. Store raw material outside the active repository so earlier answers and rubrics do not become model context.")
    w.table("How each native comparison measurement is collected", ["Measurement", "Collection rule"], [
        ["Correct supported facts, /6", "Read the saved answer against its predefined repository-specific six-fact rubric; sum supported criteria."],
        ["Incorrect or unsupported claims", "Mark distinct factual claims against original code/tests; record claim text and reason separately."],
        ["Total answer time, seconds", "Stopwatch or harness wall time from submission until final response finishes, including preparation and tool work. UI Worked timers are not automatically this measurement."],
        ["Observed tool calls", "Count visible tool invocations once; for CLI streams deduplicate step indices. Retain successful and failed counts separately."],
        ["Observed file-reading calls", "Count successful visible operations returning repository source, including content-reading commands; one call returning several files remains one call. Record searches and unique files separately when possible."],
        ["Actual reported input tokens", "Use genuine provider/CLI terminal usage once when reported. Do not sum cumulative step and terminal counters or add cache counters twice. Missing telemetry stays N/A."],
        ["Local source-context tokens", "Optional separate tokenizer count of captured source/packets; not actual provider input and not proof of all internal reads."],
        ["Outcome validity", "Record failures, permissions, task/snapshot matching, fresh delivery and unchanged source; require a matched valid pair for comparative outcomes."],
    ])
    w.table("Single-pair WITH/WITHOUT results template", ["Measurement", "WITHOUT TokenWise", "WITH TokenWise"], [
        ["Repository / task / repetition", "[INSERT]", "[INSERT matched task]"],
        ["Correct supported facts", "[INSERT score] /6", "[INSERT score] /6"],
        ["Incorrect or unsupported claims", "[INSERT count]", "[INSERT count]"],
        ["Total answer time", "[INSERT seconds]", "[INSERT seconds]"],
        ["Observed tool / file-reading calls", "[INSERT tool count / read count]", "[INSERT tool count / read count]"],
        ["Actual reported input tokens", "N/A until genuinely reported", "N/A until genuinely reported"],
        ["Run status and saved evidence", "[INSERT status / path]", "[INSERT status / path]"],
    ])
    w.table("Overall live results template for 20 repositories and 120 planned runs", ["Measurement", "WITHOUT, 60 planned runs", "WITH, 60 planned runs"], [
        ["Attempted / successful runs", "[INSERT counts] /60 planned", "[INSERT counts] /60 planned"],
        ["Valid matched pairs used for summary", "[INSERT n] /60 planned pairs", "Same matched set"],
        ["Mean correct supported facts", "[INSERT mean] /6", "[INSERT mean] /6"],
        ["Mean incorrect or unsupported claims", "[INSERT mean] per answer", "[INSERT mean] per answer"],
        ["Median total answer time", "[INSERT seconds]", "[INSERT seconds]"],
        ["Median observed tool calls", "[INSERT count]", "[INSERT count]"],
        ["Median observed file-reading calls", "[INSERT count]", "[INSERT count]"],
        ["Median actual reported input tokens", "N/A until reported", "N/A until reported"],
        ["Matched pairs with comparable reported input usage", "[INSERT n] /60 planned pairs", "Same telemetry-matched set"],
        ["Quality outcomes: WITH better / tied / worse", "[INSERT paired counts and metric]", "[INSERT paired counts and metric]"],
    ])
    repositories = json.loads(MANIFEST.read_text(encoding="utf-8"))["repositories"]
    w.table("Per-repository live result worksheet, three planned pairs each", ["Repository", "Valid pairs /3", "Mean facts /6 WO | W", "Mean claims WO | W", "Median time s WO | W", "Median reads WO | W", "Median reported input WO | W"], [
        [entry["folder"], "Pending", "Pending", "Pending", "Pending", "Pending", "N/A | N/A"]
        for entry in repositories
    ], small=True)
    w.p("WO denotes WITHOUT and W denotes WITH. The per-repository table summarizes three matched repetitions only after they exist. Failed runs are not zero scores; missing counters are not zero tokens. Compute quality means and timing/read medians on valid matched pairs, report sample count for each metric and keep failures separately. If only some pairs expose comparable usage, report that smaller denominator and do not impute the rest. An unreported read trace is N/A, while a complete observed trace with no successful reads can legitimately be zero.")
    w.code("Mean supported facts = sum of rubric scores / number of scored valid runs\nMedian paired time difference = median(time_WITH - time_WITHOUT)\nPer-pair input change (%) = 100 x (input_WITHOUT - input_WITH) / input_WITHOUT, only if input_WITHOUT > 0\nAggregate comparable input reduction (%) = 100 x (sum(input_WITHOUT) - sum(input_WITH)) / sum(input_WITHOUT)")
    w.p("Use the identical telemetry-matched pair set for both input sums. A negative time difference favors WITH on time; a positive token-change percentage indicates less reported input. Report genuine increases and ties. Input reduction is not necessarily proportional billed-cost reduction, especially when caching or pricing differ. Three repetitions are descriptive evidence, not a strong population-level significance claim.")
    w.instruction("ADD SCREENSHOT - Figure 27: After completing the real runs, capture the filled summary or validated imported comparison with the actual sample counts, supported-fact scores, time/read measurements, reported usage or N/A, and failure accounting. Caption: Observed paired WITH/WITHOUT comparison results. Leave this unfilled until measured.")
    w.h("7.15 Current Execution Status and Validity Limits")
    w.p("As of 9 October 2026, zero valid native comparison pairs are saved. One fresh CLI WITHOUT preflight on Click, using gemini-3.8-flash-high, authenticated but its repository-search command was denied. It ended with an empty answer despite terminal status SUCCESS. The captured wall time was 25.373 seconds and the terminal reported 11,696 input tokens. Those are observations of a failed preflight, not baseline effectiveness results. They are excluded from the live summary tables. The raw trace and explanation are in demonstration/comparative_study/PREFLIGHT.md and its local results directory. No fabricated example values are used in this report.")
    w.instruction("REPLACE AFTER EXECUTION - Update the preceding status paragraph with actual completed/failed/replacement counts, dates, model and evidence paths. Retain the failed preflight in the audit record rather than counting it as a successful answer or silently deleting it.")
    w.table("Threats and interpretation boundaries", ["Threat", "Required handling"], [
        ["Small purposive sample; one task per repository", "Describe exact source pins/tasks and three repetitions; do not generalize to every language or task."],
        ["Matched settings or source differ", "Reject the comparison or rerun a fresh matched pair; source snapshots must stay inside the selected workspace."],
        ["Cold start, order and cloud variability", "Alternate order, record backend state and use repeated total workflow measurements."],
        ["Incomplete/subjective answer scoring", "Freeze six-fact rubrics, retain claim reasons/citations, and request supervisor review."],
        ["Failed/denied or contaminated baseline", "Report it separately; keep native tools available and do not force WITHOUT to fail."],
        ["Local packets confused with reported usage", "Separate source counts, prepared packets and provider counters; preserve N/A."],
        ["High retrieval coverage but missing packed bodies", "Inspect actual implementation/assertion excerpts; filenames and candidate coverage are not answer quality."],
        ["Carbon predictions omit added local work", "Report estimator assumptions and excluded overhead, not measured or net emissions."],
        ["Controlled memory replay / fixture UI", "Label replay and synthetic rendering separately from native chat capture and live workflow evidence."],
    ])
    w.h("7.16 Reproduction and Final Comparative Interpretation")
    w.p("The completed component study retains cases, full source pins, metrics, summaries and the failed CPU pilot under evaluation/comparative-study. Research packet/archive files live in ignored scratch directories and may be absent in another checkout. Re-running measurements can replace artifacts; save the dated originals first. The report revision did not rerun the full software suites or the component experiment.")
    w.code(".\\.venv\\Scripts\\python.exe evaluation/comparative-study/test_protocol.py -v\n.\\.venv\\Scripts\\python.exe scripts/run_comparative_study.py --prepare-only\n.\\.venv\\Scripts\\python.exe scripts/run_comparative_study.py\n.\\.venv\\Scripts\\python.exe scripts/run_comparative_study.py --rescore\n.\\.venv\\Scripts\\python.exe scripts/summarize_comparative_study.py")
    w.p("Each line is a separate development-checkout command, not a prerequisite for installed-extension users. Preparation may download pinned public source; component measurement uses local weights and CPU, not paid Antigravity calls. Live comparison records use demonstration/comparative_study/results-template.csv and the protocol in LIVE_STUDY.md. The current capture helper supports baseline trials only and is not a completed 120-run paired harness.")
    w.p("The completed component results support bounded context management: retrieval-only packets are 93.58% smaller in aggregate than the defined all-Python baseline. Actual neural pruning adds 3.25% aggregate reduction on matched focused excerpts but loses three implementation anchors. Selected history improves candidate-file coverage in the controlled ablation but does not guarantee retained bodies. These findings remain valid within their scopes; they do not supply missing live accuracy, tool-read, timing or provider-usage values.")
    w.instruction("FINAL RESULT PARAGRAPH - After execution replace this instruction with: 'Across [n] valid matched pairs from [r] repositories, mean supported facts were [WITHOUT]/6 and [WITH]/6; median total answer times were [WITHOUT] s and [WITH] s; median observed file-reading calls were [WITHOUT] and [WITH]. Comparable reported input usage was available for [k] pairs and changed by [signed aggregate percentage or N/A]. [State failures, ties, regressions and limitations.]' Use measured values only; do not assert improvement in an unmeasured metric.")
    return w.blocks


def field_paragraph(document, instruction):
    paragraph = document.add_paragraph(style="TW Body")
    for kind in ("begin", "instruction", "separate", "end"):
        run = paragraph.add_run()
        if kind == "instruction":
            node = OxmlElement("w:instrText")
            node.set(qn("xml:space"), "preserve")
            node.text = instruction
        else:
            node = OxmlElement("w:fldChar")
            node.set(qn("w:fldCharType"), kind)
        run._r.append(node)
    paragraph._p.getparent().remove(paragraph._p)
    return paragraph._p


def main():
    source_hash = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    document = Document(SOURCE)
    body = document._element.body
    original = list(body)
    if text(original[285]) != "6. User Manual" or text(original[406]) != "7.12 Native Antigravity With and Without TokenWise":
        raise ValueError("The source structure changed; inspect the report again before applying this revision")
    configure_styles(document)
    early_fixes = {
        47: "TokenWise demonstrates an implemented approach to automatic, bounded and inspectable repository-context preparation. The reported token measurements concern locally constructed packets, not verified Antigravity provider usage. Carbon values are estimates rather than measured emissions. Controlled downstream answer-quality evaluation remains necessary to establish broader effectiveness.",
        53: "The list identifies the report's embedded figures and explicit future screenshot locations. Entries labeled pending require genuine captures after the corresponding validation. Page numbers are generated by the table of contents after document pagination.",
        158: "term_weight = ln(1 + number_of_files / (1 + content_document_frequency))\nfile_term_evidence = 5 * path_match + 3 * symbol_match + min(content_count, 4)\nfile_score = sum(term_weight * file_term_evidence for matching query terms)",
        180: "token_difference = B - P\nreduction_percent = 100 * (B - P) / B, for B > 0",
        194: "The normal TokenWise architecture does not require PostgreSQL or another relational service. Its data design concerns repository snapshots, derived caches, registration links and file ownership rather than a relational database schema.",
        221: "prefill_J = max(0, reference_prefill_J) * actual_input_tokens / 256\ndecode_J = max(0, reference_decode_J) * expected_output_tokens / 128\ntotal_J = prefill_J + decode_J\nCO2_g = total_J / 3,600,000 * intensity_g_per_kWh\nestimated_difference_g = before_CO2_g - after_CO2_g",
        257: "ADD SCREENSHOT - Figure 5: Capture genuine dated extension/backend runner summaries showing the commands, passed/discovered/skipped totals and any failures. Keep the historical table aligned with its recorded date; update it only if new runs are performed. Caption: Recorded extension and backend software verification.",
        269: "ADD SCREENSHOT - Figure 6: In the demo folder, capture real unittest discovery output and app.py execution with folder/commands visible. Caption: Demonstration application verification. Do not infer live Antigravity quality from these deterministic tests.",
        280: ".\\.venv\\Scripts\\python.exe -m unittest discover -s swe-pruner/swe-pruner/tests -v\n.\\.venv\\Scripts\\python.exe demonstration/run_checks.py\n.\\.venv\\Scripts\\python.exe evaluation/comparative-study/test_protocol.py -v",
        341: "This report contains a completed twenty-repository local component study and recorded software verification. The native WITH/WITHOUT evaluation is specified below with fill-ready templates; it remains uncompleted as of this revision. One failed CLI preflight is recorded separately, not counted as answer-quality evidence.",
        353: "The checklist specifies evidence to capture during live validation; its entries are not newly executed native checks. Use the manual's input traces and completed transport output for operational evidence. Perform any reversible source-change/cache test separately, restore the original source and rerun the demonstration tests before matched comparison trials.",
        374: "ADD SCREENSHOT - Figure 24: Capture the saved twenty-repository component-study table with the measured date, seven condition labels and totals. Caption: Recorded twenty-repository context-efficiency study. This is not a live Antigravity answer-quality table.",
    }
    for index, replacement in early_fixes.items():
        replace_text(original[index], replacement)
    for index in (257, 269, 374):
        Paragraph(original[index], document._body).style = document.styles["TW Instruction"]
    manual = build_manual(document, original)
    comparison = build_native_comparison(document)
    # Drop only old manual/comparison blocks; retain the student's technical chapters,
    # images, completed component measurements, conclusion and cover details.
    revised = original[:285] + manual + original[338:406] + comparison + original[450:]
    for element in list(body):
        body.remove(element)
    for element in revised:
        body.append(element)
    repairs = {"\u00e2\u20ac\u2122": "'", "\u00e2\u20ac\u201c": "-", "\u00e2\u20ac\u201d": "-", "\u00e2\u20ac\u00a2": "-", "\ufffd": "-"}
    for node in document._element.iter(qn("w:t")):
        if node.text:
            for broken, fixed in repairs.items():
                node.text = node.text.replace(broken, fixed)
    # Complete citations that were referenced in the original but absent from its
    # two-entry bibliography. Keep the verified supplied-paper entries intact.
    references = Writer(document)
    entries = [
        "R3. Wahid, A. B. TokenWise: Sustainable Context Optimization for Coding Agents. Project proposal, Institute of Information Technology, University of Dhaka. Supplied spl3-1442.docx.pdf. The proposal states objectives, not measured outcomes.",
        "R4. TokenWise implementation and third-party attribution, inspected 9 October 2026: vscode-extension/package.json; services/automaticSetup.ts and repositoryIndexSync.ts; swe_pruner/prune_wrapper.py, repository modules, retrieval/context_builder.py and workspace_context.py; carbon-engine training; swe_pruner/carbon_artifacts/cv_metrics.json and validation_report.json; THIRD_PARTY_NOTICES.md. Public source: https://github.com/adnan-bin-wahid/Tokenwise-updated.",
        "R5. TokenWise Test and Verification Report, tests.md; evaluation/test-validation/browser-checks.json; extension/backend test directories; demonstration/run_checks.py. Historical verification recorded 9 October 2026, distinct from live model effectiveness.",
        "R6. TokenWise Comparative Study, comparative_study.md; evaluation/comparative-study/cases.json, snapshots.json, metrics.csv, results.json, summary.json and full-file-pilot.json; scripts/run_comparative_study.py and summarize_comparative_study.py. Completed local component measurements, 9 October 2026.",
        "R7. TokenWise user and evaluation instructions: README.md, demonstation.md, demonstration2.md, validation.md, demonstration/tokenwise_demo and demonstration/comparative_study/{TASKS.md, manifest.json, LIVE_STUDY.md, PREFLIGHT.md, results-template.csv}. Live answer-quality worksheets are planned/uncollected unless backed by saved successful runs.",
        "R8. DeHalu: Agentic Hallucination Detection and Mitigation in Local CodeLLMs. Supplied Frineds_report/DeHalu_Final_Report.docx. Used only for report organization and presentation patterns, not TokenWise implementation, personal details or results.",
    ]
    for entry in entries:
        references.p(entry)
    ref_heading = next(e for e in body if text(e) == "References")
    reference_start = list(body).index(ref_heading)
    after_refs = list(body)[reference_start:]
    last_reference = next(e for e in reversed(after_refs) if text(e).startswith("R2."))
    position = list(body).index(last_reference) + 1
    for block in references.blocks:
        body.insert(position, block)
        position += 1
    # Number tables by actual caption order, not by the earlier draft's figure plan.
    captions = []
    started = False
    for element in body:
        value = text(element)
        if value == "1. Project Overview":
            started = True
        match = re.match(r"Table (?:\d+|AUTO)\. (.+)", value)
        if started and match and element.tag == qn("w:p"):
            title = match.group(1).rstrip(".")
            captions.append(title)
            replace_text(element, f"Table {len(captions)}. {title}.")
    figure_titles = [
        "TokenWise runtime architecture and processing boundaries",
        "Task-to-context activity flow and the separate overview path",
        "Repository metadata, cache and owned-file relationships",
        "Automatic prompt retrieval and asynchronous result enrichment",
        "Recorded extension/backend software verification (pending screenshot)",
        "Demonstration application verification (pending screenshot)",
        "Prerequisite and environment information",
        "Download page and VSIX installation guide",
        "Seven-stage managed backend illustration",
        "Step 3/7 private-environment setup notification",
        "Installed TokenWise extension and controls",
        "Recorded project-overview conversation with retrieval acknowledgement",
        "Command palette",
        "Pruning input-mode selection",
        "Requested threshold input",
        "Repository input trace, memory state and goal",
        "Repository token, carbon and file-provenance values",
        "Selected-file Session-class pruning result",
        "Same-chat account-lockout follow-up invocation",
        "Completed follow-up memory inspector (pending screenshot)",
        "Outgoing task-aware guidance (pending screenshot)",
        "Local automatic packet comparison (pending screenshot)",
        "Verified backend readiness (pending screenshot)",
        "Recorded twenty-repository component study (pending screenshot)",
        "Native WITHOUT answer and tools (pending successful run)",
        "Matched WITH answer and TokenWise delivery (pending successful run)",
        "Measured paired-run results (pending evaluation)",
    ]
    index_writer = Writer(document)
    index_writer.table("REMOVE", ["Figure", "Caption / status"], [[n, title] for n, title in enumerate(figure_titles, 1)])
    figure_table = index_writer.blocks[-1]
    old_figures = original[54]
    body.replace(old_figures, figure_table)
    index_writer = Writer(document)
    index_writer.table("REMOVE", ["Table", "Title"], [[n, title] for n, title in enumerate(captions, 1)])
    body.replace(original[59], index_writer.blocks[-1])
    # The original contents page was blank. Add an editable Word TOC field rather
    # than inventing page numbers before Word has laid out the revision.
    body.insert(list(body).index(original[48]) + 1, field_paragraph(document, ' TOC \\o "1-3" \\h \\z \\u '))
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    document.settings._element.append(update)
    footer = document.sections[0].footer
    if not any(p.text.strip() for p in footer.paragraphs):
        paragraph = footer.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        instruction = OxmlElement("w:instrText")
        instruction.text = " PAGE "
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        for node in (begin, instruction, end):
            run = paragraph.add_run()
            run.font.name = "Times New Roman"
            run.font.size = Pt(10)
            run._r.append(node)
        document.sections[0].different_first_page_header_footer = True
    document.core_properties.title = "TokenWise: Sustainable Context Optimization for Coding Agents"
    document.core_properties.subject = "Revised manual, screenshot interpretation and measured-comparison worksheets"
    if hashlib.sha256(SOURCE.read_bytes()).hexdigest() != source_hash:
        raise RuntimeError("The original report changed during revision; no output has been saved")
    document.save(OUTPUT)
    with ZipFile(SOURCE) as old, ZipFile(OUTPUT) as new:
        old_media = {n: hashlib.sha256(old.read(n)).hexdigest() for n in old.namelist() if n.startswith("word/media/")}
        new_media = {n: hashlib.sha256(new.read(n)).hexdigest() for n in new.namelist() if n.startswith("word/media/")}
        if old_media != new_media:
            raise RuntimeError("Embedded media did not survive intact")
    original_media = len(old_media)
    reopened = Document(OUTPUT)
    all_text = "\n".join(p.text for p in reopened.paragraphs)
    if "effectivenes\n" in all_text or any(broken in all_text for broken in repairs):
        raise RuntimeError("Uncorrected text corruption")
    NOTES.write_text(
        "# TokenWise Report Revision\n\n"
        "Revision date: 9 October 2026. Original: `TokenWise_Final_Report.docx`. "
        "Revised copy: `TokenWise_Final_Report_Updated.docx`. The original file is unchanged.\n\n"
        "## Corrected Before the Manual\n\n"
        "- Repaired broken punctuation and completed the abstract's unfinished final sentence.\n"
        "- Separated formulas and reproduction commands onto individual lines.\n"
        "- Retained cover identity, supervisor, submission date, diagrams and technical structure.\n"
        "- Reconciled figure/table lists with actual content and pending screenshot locations.\n"
        "- Restored missing R3-R8 bibliography entries used by in-text citations.\n"
        "- Added an editable table of contents and footer page-number field.\n"
        "- Historical 193 extension / 141 backend passes plus one skip / 20 demo tests remain dated records, not newly rerun results.\n\n"
        "## Revised Manual and Comparison\n\n"
        "- Preserved every embedded media part byte-for-byte and explained all 13 manual screenshots.\n"
        "- Described every visible repository token statistic, included-file row, threshold policy and carbon assumption.\n"
        "- Described selected-file input, relevance, 117-to-26 token reduction and phase-energy values.\n"
        "- Distinguished website illustrations, completed output and still-running chat invocations.\n"
        "- Added normal installation/retry, no-selection, selected-source, memory, guidance, export and cleanup instructions.\n"
        "- Kept completed local component results intact and separate from the planned live study.\n"
        "- Added a single-pair table, overall 120-run summary, 20-repository worksheet, metric definitions and formulas.\n"
        "- Left all live result fields pending; no plausible/synthetic values were inserted.\n\n"
        "## Before Submission\n\n"
        "1. Run the paired evaluation with frozen task-specific rubrics and save real evidence.\n"
        "2. Fill the tables, setting/date fields and final result paragraph; update the execution-status paragraph.\n"
        "3. Insert genuine screenshots at red ADD SCREENSHOT instructions; remove each instruction after insertion.\n"
        "4. Keep N/A for missing actual provider counters; never replace them with local-token estimates.\n"
        "5. Refresh Word's table of contents and check layout/page references after final edits.\n"
        "6. Align the abstract/conclusion with the completed live evidence without replacing the scoped historical component results.\n",
        encoding="utf-8",
    )
    print(json.dumps({"output": str(OUTPUT), "source_unchanged": hashlib.sha256(SOURCE.read_bytes()).hexdigest() == source_hash,
                      "preserved_media_parts": original_media, "tables": len(reopened.tables),
                      "figure_list_entries": len(figure_titles), "numbered_tables": len(captions)}))


if __name__ == "__main__":
    main()
